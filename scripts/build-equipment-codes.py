#!/usr/bin/env python3
"""Generuje data/equipment-codes.js z arkusza „Wypożyczalnia” (OWiT sprzęt.xlsx).

Użycie (z głównego folderu strony):
    python3 scripts/build-equipment-codes.py "/ścieżka/do/OWiT sprzęt.xlsx"

Plik służy do podpowiedzi w polu „Kod urządzenia” w panelu pracownika (katalog.html).
Strona jest publiczna, więc do pliku trafiają wyłącznie nazwy i kody sprzętu —
bez loginów, haseł, kluczy licencyjnych i uwag z arkusza.
Kody oznaczone w arkuszu jako będące „prywatnie” u kogoś są pomijane.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

SHEET = "Wypożyczalnia"
CODE_RE = re.compile(r"^[A-Z]{2,4}-\d{1,3}$")
OUT = Path(__file__).resolve().parent.parent / "data" / "equipment-codes.js"

# Nazwy pozycji w katalogu na stronie, które różnią się od nazw w arkuszu.
# Pozycje bez dopasowania dostają w panelu podpowiedź ze wszystkimi kodami.
ALIASES = {
    "Apple MacBook Pro": ["Apple Macbook Pro M2 pro", "Appple Macbook Pro M4 Max"],
    "Apple MacBook Air 15": ["Apple Macbook Air"],
    "iPhone 17 Pro": ["iPhone 17 Pro Max"],
    "Brother MFC-J2330DW": ["Brother MFC-J20330DW"],
    "ClassVR Premium VRP64/4": ["Okulary VR zestaw 4 szt"],
    "Brailliant BI20x": ["Braillant BI20x"],
    "Brailliant BI40x": ["Braillant BI40x"],
    "Prodigi Connect 12": ["Prodigy Connect 12"],
    "One Step Reader": ["OneStep reader"],
    "Clear Reader+": ["ClearReader+"],
    "NaviBelt": ["Navi Belt"],
    "Roger Clip-on Mic": ["Roger ClipOn Mic"],
    "Phonak Roger Neckloop": ["Roger Neckloop"],
    "SimplyWorks 125": [
        "Simplyworks 125 blue", "Simplyworks 125 green",
        "Simplyworks 125 red", "Simplyworks 125 yellow",
    ],
    "Mysz adaptowana - interfejs": ["Mysz adaptowana"],
    "Klawiatura Clevy z ramką": ["klawiatura Clevy"],
    "Podnóżek energetyzujący pod stopy Fellowes": ["Podnóżek energetyzujący"],
}


def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def main(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"Brak arkusza „{SHEET}” w pliku (są: {', '.join(wb.sheetnames)}).")

    groups = []
    skipped = []
    current = None
    for row in wb[SHEET].iter_rows(min_row=2, values_only=True):
        name, code = clean(row[0]), clean(row[1]).upper()
        if name:
            current = (name, [])
            groups.append(current)
        if current and CODE_RE.match(code):
            # kolumny „UWAGI” i „Na stanie”, np. „U Hanki - prywatnie”
            note = " ".join(clean(c) for c in row[7:9]).lower()
            if "prywatnie" in note:
                skipped.append(code)
                continue
            current[1].append(code)

    groups = sorted((g for g in groups if g[1]), key=lambda g: g[0].lower())

    known = {name for name, _ in groups}
    for site_name, targets in ALIASES.items():
        for t in targets:
            if t not in known:
                print(f"  UWAGA: alias „{site_name}” wskazuje na „{t}”, którego nie ma w arkuszu")

    lines = [
        "/**",
        " * data/equipment-codes.js",
        " * Kody sprzętu do podpowiedzi w polu „Kod urządzenia” (panel pracownika w katalog.html).",
        " * Plik generuje scripts/build-equipment-codes.py z arkusza OWiT sprzęt.xlsx — nie edytuj ręcznie.",
        " */",
        "window.EQUIPMENT_CODES = {",
        "  groups: [",
        ",\n".join("    " + json.dumps([name, codes], ensure_ascii=False, separators=(",", ":")) for name, codes in groups),
        "  ],",
        "  aliases: " + json.dumps(ALIASES, ensure_ascii=False, indent=2).replace("\n", "\n  ") + "",
        "};",
        "",
    ]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Zapisano {OUT} — {len(groups)} pozycji, {sum(len(c) for _, c in groups)} kodów.")
    if skipped:
        print(f"Pominięto kody oznaczone jako prywatnie ({len(skipped)}): {', '.join(skipped)}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
